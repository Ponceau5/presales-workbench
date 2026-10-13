"""GSA02 demo costs: real IO-list counts, fictional CRM prices and fees."""
import json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook
from pydantic import BaseModel, Field

STANDARDS=[{'country':'泰国','hotel':450,'allowance':180,'flight':3500,'visa':600,'logistics_rate':0.015,'logistics_fixed':3000}, {'country':'马来西亚','hotel':500,'allowance':200,'flight':3800,'visa':500,'logistics_rate':0.018,'logistics_fixed':3500}, {'country':'新加坡','hotel':900,'allowance':250,'flight':4200,'visa':400,'logistics_rate':0.02,'logistics_fixed':4000}]
class CostParameters(BaseModel):
    country:str='泰国'
    person_days:float=Field(default=40,ge=0,le=10000)
    people:int=Field(default=2,ge=1,le=100)
    daily_labor:float=Field(default=1600,ge=0,le=100000)
    hotel:float=Field(default=450,ge=0,le=100000)
    allowance:float=Field(default=180,ge=0,le=100000)
    flight:float=Field(default=3500,ge=0,le=1000000)
    visa:float=Field(default=600,ge=0,le=100000)
    logistics_rate:float=Field(default=.015,ge=0,le=1)
    logistics_fixed:float=Field(default=3000,ge=0,le=10000000)
    tax_rate:float=Field(default=.03,ge=0,le=1)
    advance:float=Field(default=200000,ge=0,le=100000000)
    months:float=Field(default=3,ge=0,le=120)
    annual_rate:float=Field(default=.05,ge=0,le=1)
    crm_multiplier:float=Field(default=1,gt=0,le=100)

def prepare(source:Path,output:Path):
    w=load_workbook(source,read_only=True,data_only=True);counts=Counter();origins={}
    for sh in w:
        header=next(sh.values,())
        if sh.title=='软件服务器规格':
            for index,row in enumerate(sh.values,1):
                if index>1 and row[0] and row[1]:
                    key=('服务器：'+str(row[0]),' / '.join(str(v) for v in row[2:5]));counts[key]+=int(row[1]);origins.setdefault(key,[]).append(f'{sh.title}!A{index}:E{index}')
        elif sh.title=='Gateway':
            for index,row in enumerate(sh.values,1):
                if index>1 and len(row)>5 and row[4] and row[5]:
                    key=('网关',str(row[5]));counts[key]+=1;origins.setdefault(key,[]).append(f'{sh.title}!F{index}')
        elif 'Equipment Type' in header and 'Equipment Name' in header:
            for index,row in enumerate(sh.values,1):
                if index>1 and len(row)>8 and row[7] and str(row[7])!='Equipment Type':
                    name=str(row[7]);model=str(row[8]) if row[8] is not None else '待确认'
                    key=(name,model);counts[key]+=1;origins.setdefault(key,[]).append(f'{sh.title}!H{index}:I{index}')
    w.close();rows=[]
    for i,((name,model),qty) in enumerate(sorted(counts.items()),1):
        unit=75000 if name.startswith('服务器：') else 12000 if name=='网关' else 800+((i*137)%22)*250
        rows.append({'sku':f'DEMO-GSA-{i:03d}','name':name,'model':model,'qty':qty,'listUnit':unit,'costUnit':unit,'quoteUnit':0,'kind':'CRM覆盖（演示）','basis':'虚构 CRM 成本单价；数量来自 IO 清单；供货范围为演示假设','source':'; '.join(origins[(name,model)][:4])})
    data={'source':source.name,'scope_note':'IO 清单包含被监控设备，并非已批准供货 BOM。本次演示假设全部列入成本范围且 CRM 全覆盖；不代表真实供货责任。','crm':rows,'standards':STANDARDS,'defaults':CostParameters().model_dump()}
    output.write_text(json.dumps(data,ensure_ascii=False,indent=2));return data

def calculate(data,p):
    crm=[{**r,'costUnit':round(r['costUnit']*p.crm_multiplier,2)} for r in data['crm']]
    materials=round(sum(r['qty']*r['costUnit'] for r in crm),2);labor=round(p.person_days*p.daily_labor,2)
    travel=round(p.person_days*(p.hotel+p.allowance)+p.people*(p.flight+p.visa),2)
    logistics=round(materials*p.logistics_rate+p.logistics_fixed,2);tax=round((materials+labor+travel+logistics)*p.tax_rate,2)
    financing=round(p.advance*p.months/12*p.annual_rate,2);total=round(materials+labor+travel+logistics+tax+financing,2)
    ti={'1:A1':'GSA02 内部成本估算（演示）','1:B4':p.country,'1:B12':f'CNY {total:,.2f}','1:F14':materials,'1:F27':0,'1:F29':p.person_days*p.hotel,'1:F30':p.person_days*p.allowance,'1:F31':p.people*p.flight,'1:F32':p.people*p.visa,'1:F38':0,'1:F39':logistics,'1:F47':financing}
    ti.update({'1:B7':0,'1:B8':0,'1:B9':0,'2:B6':f'现场调试（{p.person_days:g}人天 × {p.daily_labor:g}元/人天）','2:F6':p.person_days,'4:B5':p.country,'4:B6':'中国公司','4:B7':'DAP','4:B8':p.months,'4:B12':0,'4:B13':0,'4:B14':0,'4:B15':0,'4:B16':0})
    for r in range(6,14):
        ti[f'2:R{r}']=0;ti[f'2:T{r}']=0
        if r>6:ti[f'2:F{r}']=0
    for r in [16,17,18,*range(20,27),33,*range(35,38),40,41,42,43,44,48,49]:ti[f'1:F{r}']=0
    return {'project':'GSA02','country':p.country,'currency':'CNY','incoterm':'演示成本口径','crm':crm,'quoteRows':[], 'assumptions':['CRM 编号、价格及供货范围均为演示假设；无供应链询价项',data['scope_note'],f'调试 {p.person_days:g} 人天；{p.people} 人，每人 {p.person_days/p.people:g} 天','税费为成本基数附加预留，不是泰国实际税率或可抵扣税计算','仅核算内部成本，不使用销售定价规则'], 'templateInputs':ti,'calculatedCells':{'1:F28':labor,'1:F45':tax,'1:F46':0,'1:F50':total,'1:F19':0,'1:F15':0,'1:F34':0},'totals':{'materials':materials,'software':0,'labor':labor,'operations':travel+logistics,'tax':tax,'financing':financing,'cost':total,'hardwareQuote':0,'serviceQuote':0,'quote':0,'grossMargin':0},'parameters':p.model_dump(),'scope_note':data['scope_note'],'source':data['source'],'overrides':{},'recognizedChange':False,'demo_kind':'gsa-cost-v1','formulaOverrides':{'1:F28':f"'1.1 调试工时预估'!F14*{p.daily_labor}",'4:B26':f"(SUM('1 成本费用预估'!F14:F44)-SUM('1 成本费用预估'!F16:F17)-SUM('1 成本费用预估'!F20:F26))*{p.tax_rate}"}}

def report(d):
    p=d['parameters'];t=d['totals']
    lines=['## GSA02 内部成本表（演示）','**CRM 编号和价格为虚构演示数据，IO 清单数量为原件统计。全部项目假设由 CRM 覆盖，无供应链采购询价。**',d['scope_note'],'### CRM 匹配结果','| CRM 编号 | 设备类型 | 原清单型号 | 数量 | 成本单价 CNY | 合计 CNY |','| --- | --- | --- | ---: | ---: | ---: |']
    lines += [f"| {r['sku']} | {r['name']} | {r['model']} | {r['qty']} | {r['costUnit']:,.2f} | {r['qty']*r['costUnit']:,.2f} |" for r in d['crm']]
    lines += ['### 成本汇总','| 成本项 | 计算依据 | 金额 CNY |','| --- | --- | ---: |',f"| 设备成本 | 清单数量 × CRM演示成本 | {t['materials']:,.2f} |",f"| 调试 | {p['person_days']:g} 人天 × {p['daily_labor']:g} | {t['labor']:,.2f} |",f"| 差旅与物流 | {p['country']}演示标准 | {t['operations']:,.2f} |",f"| 税费预留 | 成本基数 × {p['tax_rate']*100:g}% | {t['tax']:,.2f} |",f"| 资金成本 | {p['advance']:g} × {p['months']:g}/12 × {p['annual_rate']*100:g}% | {t['financing']:,.2f} |",f"| **总成本** | **不含销售定价** | **{t['cost']:,.2f}** |",'### 参数与边界',*['- '+a for a in d['assumptions']]]
    return '\n\n'.join(lines[:3])+'\n\n'+'\n'.join(lines[3:])

def export_cost(d,revision):
    # Runtime export uses the adopted calculation snapshot, never current defaults.
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    w=Workbook();main=w.active;main.title="成本汇总";crm=w.create_sheet("CRM成本明细");p=w.create_sheet("参数快照")
    p.append(["参数","值"])
    for key,value in d['parameters'].items():p.append([key,value])
    param={row[0].value:row[1].coordinate for row in p.iter_rows(min_row=2)}
    ref=lambda key:"'参数快照'!"+param[key]
    crm.append(["演示CRM编号","设备类型","原件型号","数量","基础演示单价CNY","调整后单价CNY","成本合计CNY","原件来源"])
    for i,r in enumerate(d['crm'],2):
        multiplier=d['parameters']['crm_multiplier']
        crm.append([r['sku'],r['name'],r['model'],r['qty'],round(r['costUnit']/multiplier,2),f"=E{i}*{ref('crm_multiplier')}",f"=D{i}*F{i}",r['source']])
    main.append([f"GSA02 内部成本表 V{revision}（虚构价格，仅演示）","金额 CNY","依据"])
    main.append(["设备成本",f"=SUM('CRM成本明细'!G2:G{len(d['crm'])+1})","IO清单数量×虚构CRM成本价"])
    main.append(["调试成本",f"={ref('person_days')}*{ref('daily_labor')}","人天×单价"])
    main.append(["差旅成本",f"={ref('person_days')}*({ref('hotel')}+{ref('allowance')})+{ref('people')}*({ref('flight')}+{ref('visa')})","住宿补贴+机票签证"])
    main.append(["物流成本",f"=B2*{ref('logistics_rate')}+{ref('logistics_fixed')}","设备成本比例+固定费用"])
    main.append(["税费预留",f"=SUM(B2:B5)*{ref('tax_rate')}","演示成本附加预留，非法定税率"])
    main.append(["资金成本",f"={ref('advance')}*{ref('months')}/12*{ref('annual_rate')}","垫资×期间×年率"])
    main.append(["总成本","=SUM(B2:B7)","不含销售定价"])
    main.append([d['scope_note']]);main.append(["本文件保存采纳时的清单、CRM、参数快照；实际商务使用须重新核对供货范围和费用口径。"])
    for sh in w:
        sh.freeze_panes="A2"
        for cell in sh[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='217F64')
        for col in sh.columns:
            letter=col[0].column_letter;sh.column_dimensions[letter].width=24
        for row in sh.iter_rows():
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    main.column_dimensions['A'].width=65;main.column_dimensions['C'].width=52;crm.column_dimensions['H'].width=55
    for row in main.iter_rows(min_row=2,max_row=8):row[1].number_format='#,##0.00'
    stream=BytesIO();w.save(stream);return stream.getvalue()

def recalculate_template(d, edits):
    import copy
    d=copy.deepcopy(d);inputs=d['templateInputs'];inputs.update(edits)
    def num(key):
        value=inputs.get(key,0)
        if value=='':return 0.0
        return float(value)
    calc=d['calculatedCells'];calc['1:F15']=num('1:F16')+num('1:F17');calc['1:F19']=sum(num(f'1:F{r}') for r in range(20,27))
    days=sum(num(f'2:F{r}') for r in range(6,14));labor=days*d['parameters']['daily_labor'];after=sum(num(f'2:R{r}')*num(f'2:T{r}') for r in range(6,14))
    calc.update({'1:F28':labor,'1:F34':after,'2:F14':days,'2:U14':after})
    for r in range(6,14):calc[f'2:U{r}']=num(f'2:R{r}')*num(f'2:T{r}')
    direct=num('1:F14')+calc['1:F15']+num('1:F18')+calc['1:F19']+num('1:F27')+labor+sum(num(f'1:F{r}') for r in range(29,34))+after+sum(num(f'1:F{r}') for r in range(35,45))
    tax=round(direct*d['parameters']['tax_rate'],2);calc.update({'1:F45':tax,'1:F46':0,'4:B20':0,'4:B21':0,'4:B22':0,'4:B23':0,'4:B24':0,'4:B25':0,'4:B26':tax,'4:B28':tax,'4:B30':0,'4:B31':tax})
    cost=round(direct+tax+sum(num(f'1:F{r}') for r in [47,48,49]),2);calc['1:F50']=cost;inputs['1:B12']=f'CNY {cost:,.2f}'
    d['totals'].update({'materials':num('1:F14')+calc['1:F15']+num('1:F18')+calc['1:F19'],'software':num('1:F27'),'labor':labor+sum(num(f'1:F{r}') for r in range(29,34))+after,'operations':sum(num(f'1:F{r}') for r in range(35,45)),'tax':tax,'financing':sum(num(f'1:F{r}') for r in [47,48,49]),'cost':cost})
    d['templateEdits']={**d.get('templateEdits',{}),**edits};return d

def export_original_template(d, revision):
    # Patch values/formula caches in the original zip; preserve all styles, merges,
    # widths, print settings, drawings and seven original worksheet names.
    import re, zipfile, html
    from io import BytesIO
    template=Path(__file__).resolve().parents[1]/'public/overseas-cost-template.xlsx'
    d=recalculate_template(d,{})
    values={**d['templateInputs'],'1:A1':f'GSA02 内部成本表 V{revision}（演示）'}
    formulas=d.get('formulaOverrides',{});cache=d['calculatedCells'];out=BytesIO()
    with zipfile.ZipFile(template) as source,zipfile.ZipFile(out,'w',zipfile.ZIP_STORED) as dest:
        for entry in source.infolist():
            raw=source.read(entry.filename);match=re.fullmatch(r'xl/worksheets/sheet([1-7]).xml',entry.filename)
            if match:
                idx=match[1];xml=raw.decode()
                def cell(m):
                    old=m.group(0);ref=m.group(1);key=idx+':'+ref
                    if key not in values and key not in formulas and key not in cache:return old
                    start=re.match(r'<x:c\b[^>]*',old).group(0).rstrip('/')
                    start=re.sub(r' t="[^"]*"','',start)
                    existing=re.search(r'<x:f[^>]*>(.*?)</x:f>',old,re.S)
                    if key in formulas or existing:
                        formula=formulas.get(key,html.unescape(existing.group(1)) if existing else '')
                        cached=cache.get(key)
                        tag=f'<x:v>{cached}</x:v>' if isinstance(cached,(int,float)) else ''
                        return start+'><x:f>'+html.escape(formula,quote=False)+'</x:f>'+tag+'</x:c>'
                    value=values.get(key,cache.get(key,''))
                    if isinstance(value,(int,float)):return start+f'><x:v>{value}</x:v></x:c>'
                    return start+' t="inlineStr"><x:is><x:t xml:space="preserve">'+html.escape(str(value))+'</x:t></x:is></x:c>'
                xml=re.sub(r'<x:c\b[^>]*\br="([A-Z]+\d+)"[^>]*?(?:/>|>.*?</x:c>)',cell,xml,flags=re.S);raw=xml.encode()
            if entry.filename=='xl/workbook.xml':
                xml=raw.decode();xml=re.sub(r'<x:calcPr\b[^>]*/>','',xml);raw=xml.replace('</x:workbook>','<x:calcPr calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/></x:workbook>').encode()
            dest.writestr(entry,raw)
    return out.getvalue()
