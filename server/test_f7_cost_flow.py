import os,tempfile,unittest,io
from uuid import uuid4
from openpyxl import Workbook,load_workbook
TMP=tempfile.TemporaryDirectory(dir=os.environ.get('PRESALES_TEST_TMP'));os.environ['PRESALES_DATA_DIR']=TMP.name
from app import app
from fastapi.testclient import TestClient
class ProcurementTests(unittest.TestCase):
 def test_export_and_match(self):
  c=TestClient(app);auth=c.post('/api/session',json={'username':'pm','password':'demo2026'}).json();h={'Authorization':'Bearer '+auth['token']};p='TEST-'+uuid4().hex[:6]
  c.post('/api/access/projects',headers=h,json={'id':p,'name':p});base='/api/projects/'+p+'/workspaces/F7/cost-flow'
  def upload(title,headers,row):
   w=Workbook();s=w.active;s.append(headers);s.append(row);b=io.BytesIO();w.save(b)
   r=c.post('/api/projects/'+p+'/documents',headers=h,data={'title':title,'version':'V1'},files={'file':('input.xlsx',b.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')});self.assertEqual(r.status_code,200,r.text);return r.json()['id']
  equipment=upload('配置清单',['料号','设备名称','规格型号','数量','采购类型'],['A01','服务器','M1',2,'外购'])
  self.assertEqual(c.post(base+'/input',headers=h,json={'slot':'equipment','document_id':equipment}).status_code,200)
  status=c.get(base,headers=h).json();self.assertEqual(status['bom'][0]['sku'],'A01')
  inquiry=c.post(base+'/enquiry',headers=h,json={'rows':status['bom'],'expected_revision':0});self.assertEqual(inquiry.status_code,200,inquiry.text)
  data=c.get(base+'/enquiry/'+inquiry.json()['id']+'/download',headers=h);self.assertEqual(data.status_code,200)
  w=load_workbook(io.BytesIO(data.content));self.assertEqual(w.active['A2'].value,'A01');self.assertIsNone(w.active['E2'].value)
  self.assertEqual(c.post(base+'/enquiry',headers=h,json={'rows':status['bom'],'expected_revision':0}).status_code,409)
  prices=upload('采购回价',['料号','设备名称','规格型号','数量','采购单价'],['A01','服务器','M1',2,15000])
  c.post(base+'/input',headers=h,json={'slot':'purchase','document_id':prices});status=c.get(base,headers=h).json();self.assertEqual(status['matched'],['A01']);self.assertEqual(status['missing'],[])
