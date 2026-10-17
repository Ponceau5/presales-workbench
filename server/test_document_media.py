"""Attachment storage, actual image transport and original-source access boundaries."""
import os,tempfile,io,unittest,json
from pathlib import Path
from unittest.mock import patch
from PIL import Image
TEMP=tempfile.TemporaryDirectory(prefix='media-',dir=os.environ.get('PRESALES_TEST_TMP'));os.environ['PRESALES_DATA_DIR']=TEMP.name
from fastapi.testclient import TestClient
from app import app,db
from model_gateway import ModelConfig
class MediaTests(unittest.TestCase):
 def setUp(self):
  self.client=TestClient(app);r=self.client.post('/api/session',json={'username':'pm','password':'demo2026'});self.headers={'Authorization':'Bearer '+r.json()['token']}
 def upload(self,name,data):
  r=self.client.post('/api/projects/GSA02/documents',headers=self.headers,data={'title':name,'version':'test'},files={'file':(name,data)})
  self.assertEqual(r.status_code,200,r.text);return r.json()
 def png(self):
  out=io.BytesIO();Image.new('RGB',(32,32),'red').save(out,format='PNG');return out.getvalue()
 def test_unknown_file_preserved_without_claiming_text(self):
  uploaded=self.upload('设备图.dwg',b'\x00\xffbinary drawing')
  self.assertEqual(uploaded['page_count'],0);self.assertEqual(uploaded['parse_status'],'unparsed')
  self.assertEqual(self.client.get('/api/documents/'+uploaded['id']+'/download',headers=self.headers).content,b'\x00\xffbinary drawing')
 def test_image_is_sent_to_model_as_image_not_filename(self):
  uploaded=self.upload('截图.png',self.png());identifier=uploaded['id'];self.assertTrue(uploaded['is_image'])
  with patch('workspace_router.model_config',return_value=ModelConfig('zhipu','test','https://example.invalid','mock')),patch('workspace_router.request_agent_stream',return_value=iter(['原图如下'])) as model:
   r=self.client.post('/api/projects/GSA02/workspaces/F3/ask-stream',headers=self.headers,json={'turn_id':'image-read-test','task_id':'deliverable','question':'图里是什么','attachment_ids':[identifier]})
  self.assertIn('"type": "done"',r.text)
  content=model.call_args.args[0][1]['content'];self.assertIsInstance(content,list);self.assertTrue(content[1]['image_url']['url'].startswith('data:image/png;base64,'))
  figure=self.client.get(f'/api/documents/{identifier}/figures/1',headers=self.headers);self.assertEqual(figure.status_code,200);self.assertEqual(Image.open(io.BytesIO(figure.content)).getpixel((0,0)),(255,0,0))
  self.assertEqual(self.client.get(f'/api/documents/{identifier}/figures/1').status_code,401)
 def test_embedded_word_figure_and_cross_project_denial(self):
  from docx import Document
  doc=Document();doc.add_paragraph('设备接线参考');doc.add_picture(io.BytesIO(self.png()));out=io.BytesIO();doc.save(out)
  identifier=self.upload('参考.docx',out.getvalue())['id']
  self.assertEqual(self.client.get(f'/api/documents/{identifier}/figures/1',headers=self.headers).status_code,200)
  user=self.client.post('/api/session',json={'username':'sales','password':'demo2026'}).json();headers={'Authorization':'Bearer '+user['token']}
  with db() as c:c.execute("DELETE FROM project_owners WHERE user_id='sales' AND project_id='GSA02'")
  self.assertEqual(self.client.get(f'/api/documents/{identifier}/figures/1',headers=headers).status_code,403)
 def test_other_project_attachment_rejected_before_model(self):
  uploaded=self.upload('边界.txt',b'project only')
  r=self.client.post('/api/projects/RCJM1/workspaces/F3/ask-stream',headers=self.headers,json={'turn_id':'cross-project-test','question':'读附件','attachment_ids':[uploaded['id']]})
  self.assertEqual(r.status_code,403)
