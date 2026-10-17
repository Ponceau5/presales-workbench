"""F7 procurement handoff: immutable enquiry rows, exact return matching and source readiness."""
import json,io,re
from uuid import uuid4
from fastapi import APIRouter,Depends,HTTPException
from fastapi.responses import Response
from pydantic import BaseModel,Field
from access_control import stamp
from object_storage import materialize
from openpyxl import load_workbook,Workbook
class Pick(BaseModel):
 slot:str
 document_id:str
class Enquiry(BaseModel):
 rows:list[dict]=Field(min_length=1,max_length=1000)
 expected_revision:int=Field(ge=0)
def records(document):
 if not document:return []
 with materialize(document['path']) as path:
  if path.suffix.lower() not in {'.xls','.xlsx','.xlsm'}:return []
  legacy=path.suffix.lower()=='.xls'
  if legacy:
   import xlrd
   w=xlrd.open_workbook(str(path),on_demand=True)
   sheets=((sh.name,(sh.row_values(i) for i in range(sh.nrows))) for sh in w.sheets())
  else:
   w=load_workbook(path,read_only=True,data_only=True)
   sheets=((sh.title,sh.iter_rows(values_only=True)) for sh in w)
  result=[]
  aliases={'sku':['料号','物料号','设备编号','CRM编号','编号','SKU'],'name':['设备名称','物料名称','名称','设备类型','设备名 Device name'],'model':['型号','规格型号','规格','Gateway Model','规格型号 Spec'],'qty':['数量','台数','Qty','数量 QTY'],'unit_cost':['成本单价','采购单价','单价','回价'],'category':['采购类型','供货类型','成本类型'],'currency':['币种','货币']}
  try:
   for sheet_name,rows in sheets:
    mapping=None
    for cells in rows:
     if mapping is None:
      headers=[str(v or '').strip() for v in cells];mapping={k:next((i for i,h in enumerate(headers) if h in options),None) for k,options in aliases.items()}
      if mapping['name'] is None or mapping['qty'] is None:mapping=None
      continue
     row={k:cells[i] if i is not None and i<len(cells) else None for k,i in mapping.items()}
     if not row['name'] or not isinstance(row['qty'],(int,float)) or row['qty']<=0:continue
     row.update(currency=str(row.get('currency') or 'CNY'),sku=str(row['sku'] or ''),name=str(row['name']),model=str(row['model'] or ''),category=str(row['category'] or '待确认'));result.append(row)
     if len(result)>1000:raise HTTPException(413,'清单超过1000行，请拆分')
  finally:
   if legacy:w.release_resources()
   else:w.close()
 return result

def install(app,db,actor):
 with db() as c:
  c.execute('CREATE TABLE IF NOT EXISTS f7_cost_inputs(project_id TEXT NOT NULL,slot TEXT NOT NULL,document_id TEXT NOT NULL,PRIMARY KEY(project_id,slot))')
  c.execute('CREATE TABLE IF NOT EXISTS f7_enquiries(id TEXT PRIMARY KEY,project_id TEXT NOT NULL,revision INTEGER NOT NULL,source_id TEXT NOT NULL,rows TEXT NOT NULL,created_at TEXT NOT NULL,UNIQUE(project_id,revision))')
 base='/api/projects/{project_id}/workspaces/F7/cost-flow';router=APIRouter(prefix=base)
 names={'equipment':'配置清单','purchase':'采购回价','software':'软件成本','labor':'工时配置表','maintenance':'维保驻场费用表','logistics':'运输清关','tax':'税费依据','cash':'付款计划'}
 patterns={'equipment':'解决方案.*清单|配置清单|设备BOM','purchase':'采购.*成本|采购.*回价|回价|供应商报价','software':'软件.*成本|定制.*成本','labor':'工时','maintenance':'驻场|维保','logistics':'物流|运输|清关','tax':'税费.*测算|税费.*依据','cash':'现金流|付款计划'}
 def inputs(c,project):
  docs=[dict(r) for r in c.execute('SELECT id,title,path,version,created_at FROM documents WHERE project_id=? ORDER BY created_at DESC',(project,))]
  chosen={r['slot']:r['document_id'] for r in c.execute('SELECT * FROM f7_cost_inputs WHERE project_id=?',(project,))}
  return {key:next((d for d in docs if d['id']==chosen.get(key)),None) or next((d for d in docs if re.search(patterns[key],d['title'])),None) for key in names},docs
 @router.get('')
 def get(project_id:str,user:str=Depends(actor)):
  with db() as c:
   selected,docs=inputs(c,project_id);enquiry=c.execute('SELECT * FROM f7_enquiries WHERE project_id=? ORDER BY revision DESC LIMIT 1',(project_id,)).fetchone()
  bom=records(selected['equipment']);prices=records(selected['purchase']);price_index={}
  for row in prices:price_index.setdefault(row['sku'],[]).append(row)
  matched=[];missing=[]
  expected=json.loads(enquiry['rows']) if enquiry else [r for r in bom if re.search('外购|采购',r['category'])]
  for row in expected:
   candidates=price_index.get(row['sku'],[]) if row['sku'] else []
   ok=len(candidates)==1 and candidates[0]['model']==row['model'] and candidates[0]['qty']==row['qty'] and isinstance(candidates[0]['unit_cost'],(int,float)) and candidates[0]['unit_cost']>=0 and candidates[0]['currency']=='CNY'
   (matched if ok else missing).append(row['sku'] or row['name'])
  return {'inputs':{k:{'name':names[k],'document':{a:d[a] for a in ('id','title','version')} if d else None} for k,d in selected.items()},'documents':[{k:d[k] for k in ('id','title','version')} for d in docs],'bom':bom,'enquiry_revision':enquiry['revision'] if enquiry else 0,'enquiry_id':enquiry['id'] if enquiry else None,'enquiry_stale':bool(enquiry and (not selected['equipment'] or enquiry['source_id']!=selected['equipment']['id'])),'matched':matched,'missing':missing,'purchase_required':not bool(bom) or bool(expected) or any(r['category']=='待确认' for r in bom)}
 @router.post('/input')
 def pick(project_id:str,payload:Pick,user:str=Depends(actor)):
  if payload.slot not in names:raise HTTPException(422,'输入类别无效')
  with db() as c:
   if not c.execute('SELECT 1 FROM documents WHERE id=? AND project_id=?',(payload.document_id,project_id)).fetchone():raise HTTPException(404,'资料不属于当前项目')
   c.execute('INSERT INTO f7_cost_inputs VALUES (?,?,?) ON CONFLICT(project_id,slot) DO UPDATE SET document_id=excluded.document_id',(project_id,payload.slot,payload.document_id))
  return {'saved':True}
 @router.post('/enquiry')
 def create(project_id:str,payload:Enquiry,user:str=Depends(actor)):
  with db() as c:
   c.execute('BEGIN IMMEDIATE');selected,_=inputs(c,project_id);revision=c.execute('SELECT COALESCE(MAX(revision),0) FROM f7_enquiries WHERE project_id=?',(project_id,)).fetchone()[0]
   if revision!=payload.expected_revision:raise HTTPException(409,'询价清单已更新，请刷新')
   if not selected['equipment']:raise HTTPException(409,'请先选择解决方案配置清单')
   source=records(selected['equipment']);allowed={(r['sku'],r['name'],r['model'],r['qty']) for r in source};seen=set()
   for r in payload.rows:
    key=(r.get('sku'),r.get('name'),r.get('model'),r.get('qty'))
    if key not in allowed:raise HTTPException(422,'询价项目与所选配置清单不一致')
    if not r['sku'] or r['sku'] in seen:raise HTTPException(422,'请在源清单中补齐唯一料号后导出询价')
    seen.add(r['sku'])
   identifier=str(uuid4());c.execute('INSERT INTO f7_enquiries VALUES (?,?,?,?,?,?)',(identifier,project_id,revision+1,selected['equipment']['id'],json.dumps(payload.rows,ensure_ascii=False),stamp()))
  return {'id':identifier,'revision':revision+1}
 @router.get('/enquiry/{identifier}/download')
 def download(project_id:str,identifier:str,user:str=Depends(actor)):
  with db() as c:r=c.execute('SELECT * FROM f7_enquiries WHERE id=? AND project_id=?',(identifier,project_id)).fetchone()
  if not r:raise HTTPException(404,'询价清单不存在')
  w=Workbook();s=w.active;s.title='采购询价清单';s.append(['料号','设备名称','规格型号','数量','采购单价','币种','回价有效期','采购备注'])
  for row in json.loads(r['rows']):s.append([row['sku'],row['name'],row['model'],row['qty'],None,'CNY',None,None])
  s.freeze_panes='A2'
  for col,width in [('A',22),('B',28),('C',30),('D',12),('E',16),('F',12),('G',20),('H',32)]:s.column_dimensions[col].width=width
  out=io.BytesIO();w.save(out);return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="procurement-enquiry.xlsx"'})
 app.include_router(router)

def export_cost_draft(snapshot,project,revision):
 """Fill a copy of the original seven-sheet template; preserve its package/styles."""
 import zipfile,xml.etree.ElementTree as ET
 from pathlib import Path
 ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'};tag='{'+ns['s']+'}'
 source=Path(__file__).resolve().parents[1]/'public/overseas-cost-template.xlsx';out=io.BytesIO()
 totals={20:0,21:0,22:0,23:0,24:0,25:0,26:0};used=set()
 for r in snapshot['known_rows']:
  name=r['name'];row=20 if '服务器' in name else 21 if re.search('交换机|网关',name) else 22 if re.search('控制柜|PLC',name) else 23 if '流量计' in name else 24 if 'ICT' in name else 25 if re.search('安防|视频|门禁',name) else 26
  totals[row]+=r['qty']*r['unit_cost'];used.add(row)
 values={'A1':f'{project} 成本估算表 V{revision}（待补草稿）','G50':'本表含已匹配采购成本；其他未确认成本待补，当前合计不代表完整项目成本。'}
 values.update({'F'+str(row):round(totals[row],2) for row in used})
 with zipfile.ZipFile(source) as src,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dst:
  for info in src.infolist():
   data=src.read(info.filename)
   if info.filename.startswith('xl/worksheets/sheet') and info.filename.endswith('.xml'):
    root=ET.fromstring(data)
    for cell in root.findall('.//s:c',ns):
     ref=cell.get('r');formula=cell.find('s:f',ns)
     if formula is not None:
      value=cell.find('s:v',ns)
      if value is not None:cell.remove(value)
     if info.filename=='xl/worksheets/sheet1.xml' and (ref in values or (ref and re.fullmatch(r'F(?:1[4-9]|[2-4][0-9])',ref) and formula is None)):
      for child in list(cell):cell.remove(child)
      cell.attrib.pop('t',None)
      if ref in values:
       value=values[ref]
       if isinstance(value,str):cell.set('t','inlineStr');ET.SubElement(ET.SubElement(cell,tag+'is'),tag+'t').text=value
       else:ET.SubElement(cell,tag+'v').text=str(value)
    data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
   elif info.filename=='xl/workbook.xml':
    root=ET.fromstring(data);calc=root.find('s:calcPr',ns)
    if calc is None:calc=ET.SubElement(root,tag+'calcPr')
    calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1');data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
   dst.writestr(info,data)
 return out.getvalue()
