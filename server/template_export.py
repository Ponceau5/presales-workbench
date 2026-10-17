"""Fill conventional table/placeholder templates; preserve source files and formulas."""
from io import BytesIO
from copy import copy,deepcopy
from pathlib import Path
import re
from fastapi import HTTPException
from file_formats import blocks

def key(value):return re.sub(r'[\s*_：:（）()]','',str(value or '')).lower()
def tables(text):return [rows for kind,rows in blocks(text) if kind=='table' and len(rows)>1]
def safe(value):
 value=re.sub(r'\*\*','',str(value)).strip()
 try:return float(value.replace(',','')) if '.' in value else int(value.replace(',',''))
 except ValueError:return "'"+value if value.startswith(('=','+','-','@')) else value

def export_template(text,kind,path):
 output=BytesIO();data=tables(text)
 if kind=='xlsx':
  from openpyxl import load_workbook
  from openpyxl.cell.cell import MergedCell
  from openpyxl.formula.translate import Translator
  book=load_workbook(path)
  filled=False
  for rows in data:
   source={key(h):i for i,h in enumerate(rows[0])};candidates=[]
   for sheet in book:
    for cells in sheet.iter_rows(min_row=1,max_row=min(30,sheet.max_row),max_col=min(80,sheet.max_column)):
     headers={cell.column:key(cell.value) for cell in cells if cell.value is not None}
     if len(headers)>=2 and set(headers.values())<=set(source):candidates.append((sheet,cells[0].row,headers))
   if not candidates:continue
   named=[v for v in candidates if v[0].title in text]
   if len(candidates)>1 and len(named)!=1:raise HTTPException(422,'模板中有多个相同表头，请在成果中明确工作表名称')
   sheet,header,headers=(named or candidates)[0]
   sample=[copy(sheet.cell(header+1,col)) for col in range(1,sheet.max_column+1)]
   end=header+1
   while end<=sheet.max_row:
    values=[sheet.cell(end,col).value for col in headers]
    if any(str(v or '').strip() in {'合计','总计','小计'} for v in values):break
    if end>header+1 and all(v is None for v in values):break
    end+=1
   for index,row in enumerate(rows[1:],header+1):
    if index>=end and end<=sheet.max_row:sheet.insert_rows(index);end+=1
    for col,field in headers.items():
     cell=sheet.cell(index,col)
     if isinstance(cell,MergedCell):continue
     original=sample[col-1]
     if not cell.has_style and original.has_style:cell._style=copy(original._style)
     if original.data_type=='f':cell.value=Translator(original.value,origin=original.coordinate).translate_formula(cell.coordinate)
     else:cell.value=safe(row[source[field]]) if source[field]<len(row) else None
   for index in range(header+len(rows),end):
    for col in headers:
     cell=sheet.cell(index,col)
     if not isinstance(cell,MergedCell) and cell.data_type!='f':cell.value=None
   filled=True
  if not filled:raise HTTPException(422,'成果表头与所选模板不一致，请按模板列名重新生成后导出')
  book.save(output);book.close()
  return output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
 if kind=='docx':
  from docx import Document
  document=Document(path);filled=False
  for table in document.tables:
   headers=[key(c.text) for c in table.rows[0].cells]
   rows=next((r for r in data if len(headers)>=2 and set(headers)<=set(map(key,r[0]))),None)
   if not rows:continue
   source={key(h):i for i,h in enumerate(rows[0])}
   original=deepcopy(table.rows[1]._tr) if len(table.rows)>1 else None
   while len(table.rows)<len(rows):
    if original is not None:table._tbl.append(deepcopy(original))
    else:table.add_row()
   for index,row in enumerate(rows[1:],1):
    for col,field in enumerate(headers):table.rows[index].cells[col].text=re.sub(r'\*\*','',row[source[field]]) if source[field]<len(row) else ''
   for row in list(table.rows)[len(rows):]:table._tbl.remove(row._tr)
   filled=True
  values={key(r[0]):r[1] for rows in data for r in rows[1:] if len(r)>=2}
  paragraphs=list(document.paragraphs)+[p for t in document.tables for row in t.rows for c in row.cells for p in c.paragraphs]
  for paragraph in paragraphs:
   if re.search(r'\{\{[^{}]+\}\}',paragraph.text):
    value=re.sub(r'\{\{([^{}]+)\}\}',lambda m:values.get(key(m[1]),'待确认'),paragraph.text)
    if paragraph.runs:
     paragraph.runs[0].text=value
     for run in paragraph.runs[1:]:run.text=''
    else:paragraph.add_run(value)
    filled=True
  if not filled:raise HTTPException(422,'Word 模板未找到匹配表头或 {{字段}} 占位符，请核对模板与输出字段')
  document.save(output)
  return output.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
 raise HTTPException(422,'仅 Word 和 Excel 模板支持原格式填充')
