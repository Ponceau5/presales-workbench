"""Registration, guide isolation, editable recipes and real template exports."""
import os,tempfile,unittest,copy
from pathlib import Path
from uuid import uuid4
from io import BytesIO
ROOT=tempfile.TemporaryDirectory(dir=os.getenv('PRESALES_TEST_TMP'))
os.environ['PRESALES_DATA_DIR']=ROOT.name
from fastapi.testclient import TestClient
from app import app
from template_export import export_template
class SkillMaterialsTests(unittest.TestCase):
 def setUp(self):
  self.c=TestClient(app)
  token=self.c.post('/api/session',json={'username':'pm','password':'demo2026'}).json()['token'];self.h={'Authorization':'Bearer '+token};self.p='T-'+uuid4().hex[:8]
  self.c.post('/api/access/projects',headers=self.h,json={'id':self.p,'name':self.p})
 def upload(self,body=b'project device list',project=None):
  r=self.c.post('/api/projects/'+(project or self.p)+'/documents',headers=self.h,data={'title':'设备清单','version':'V1'},files={'file':('list.txt',body,'text/plain')});self.assertEqual(r.status_code,200,r.text);return r.json()
 def test_digest_dedup_and_template_discovery(self):
  first=self.upload();second=self.upload();self.assertEqual(first['id'],second['id']);self.assertTrue(second['duplicate'])
  base='/api/projects/'+self.p
  r=self.c.post(base+'/documents/'+first['id']+'/register',headers=self.h,json={'mode':'template'});self.assertEqual(r.status_code,200,r.text)
  catalog=self.c.get(base+'/workspaces/F5/materials',headers=self.h);self.assertEqual(catalog.status_code,200,catalog.text);self.assertEqual(catalog.json()['templates'][0]['id'],first['id'])
  rows=self.c.get(base+'/documents',headers=self.h).json();self.assertEqual(rows[0]['category'],'template')
  reuse=self.c.post(base+'/documents/'+first['id']+'/register',headers=self.h,json={'mode':'auto'}).json();self.assertTrue(reuse['reused'])
 def test_temporary_excluded_and_cross_project_digest_isolated(self):
  doc=self.upload();base='/api/projects/'+self.p
  self.c.post(base+'/documents/'+doc['id']+'/register',headers=self.h,json={'mode':'auto','usage':'仅本次，不要保存'})
  self.assertEqual(self.c.get(base+'/documents',headers=self.h).json(),[])
  q='T-'+uuid4().hex[:8];self.c.post('/api/access/projects',headers=self.h,json={'id':q,'name':q})
  other=self.upload(project=q);self.assertNotEqual(other['id'],doc['id'])
 def test_rename_delete_and_empty_workflow(self):
  base='/api/projects/'+self.p+'/workspaces/F5/flows';cfg=self.c.get(base,headers=self.h).json();definition=copy.deepcopy(cfg['active']['definition'])
  self.assertEqual(len(definition['tasks']),5);definition['tasks'][0]['name']='我的摸排';definition['tasks'].pop()
  body={'expected_revision':0,'base_version_id':cfg['active']['id'],'definition':definition}
  saved=self.c.post(base+'/draft',headers=self.h,json=body);self.assertEqual(saved.status_code,200,saved.text)
  after=self.c.get(base,headers=self.h).json()['draft'];self.assertEqual(after['definition']['tasks'][0]['name'],'我的摸排');self.assertEqual(len(after['definition']['tasks']),4)
  body.update(expected_revision=after['revision'],definition={**definition,'tasks':[]});self.assertEqual(self.c.post(base+'/draft',headers=self.h,json=body).status_code,200)
  from workflow_studio import instructions
  self.assertEqual(instructions({'definition':{'tasks':[]}},None,''),'')
 def test_guides_persist_per_user(self):
  name='guide_'+uuid4().hex[:8];self.c.post('/api/access/users',headers=self.h,json={'username':name,'display_name':name,'role':'sales','password':'testpass2026'})
  token=self.c.post('/api/session',json={'username':name,'password':'testpass2026'}).json()['token'];h={'Authorization':'Bearer '+token}
  self.assertFalse(self.c.get('/api/me/guide',headers=h).json()['seen']);self.c.post('/api/me/guide/seen',headers=h);self.assertTrue(self.c.get('/api/me/guide',headers=h).json()['seen'])
 def test_excel_template_retains_formula_and_styles(self):
  from openpyxl import Workbook,load_workbook
  from openpyxl.styles import PatternFill
  path=Path(ROOT.name)/'template.xlsx';w=Workbook();s=w.active;s.append(['设备','数量','单价','金额']);s.append(['样例',99,99,'=B2*C2']);s['A2'].fill=PatternFill('solid',fgColor='FFFF00');w.save(path)
  data,_=export_template('|设备|数量|单价|金额|\n|---|---|---|---|\n|网关|2|30|60|\n|服务器|1|50|50|','xlsx',path)
  s=load_workbook(BytesIO(data)).active;self.assertEqual(s['A2'].value,'网关');self.assertEqual(s['D3'].value,'=B3*C3');self.assertEqual(s['A3'].fill.fgColor.rgb,'00FFFF00')
 def test_registered_template_is_in_model_context(self):
  from unittest.mock import patch
  from model_gateway import ModelConfig
  doc=self.upload('|设备|数量|\n|---|---|'.encode())
  self.c.post('/api/projects/'+self.p+'/documents/'+doc['id']+'/register',headers=self.h,json={'mode':'template'})
  self.upload(b'Confirmed project input: gateways 2')
  with patch('workspace_router.model_config',return_value=ModelConfig('zhipu','test','https://example.invalid','test')),patch('workspace_router.request_agent_stream',return_value=iter(['请确认设备数量'])) as model:
   r=self.c.post('/api/projects/'+self.p+'/workspaces/F5/ask-stream',headers=self.h,json={'question':'按设备清单模板输出功能清单','task_id':'functions','turn_id':uuid4().hex})
  self.assertEqual(r.status_code,200,r.text);self.assertNotIn('"type": "error"',r.text)
  snapshots=model.call_args.kwargs['sources'];self.assertTrue(any(v.get('id')==doc['id'] and v.get('purpose')=='output' for v in snapshots))
 def test_server_storage_tool_uses_explicit_project_scale(self):
  from agent_engine import snapshot_tools
  tool=next(t for t in snapshot_tools('pm',self.p,[],{},[],stage='F5') if t.name=='calculate_server_resources')
  import json
  result=json.loads(tool.invoke({'point_count':1000,'sampling_seconds':10,'retention_years':1,'device_count':10,'cache_days':1,'backup_copies':1,'raid_usable_ratio':0.5,'headroom':0.2,'inputs_confirmed':True}))
  self.assertEqual(result['points']['high'],1000)
 def test_word_template_table_and_placeholders(self):
  from docx import Document
  path=Path(ROOT.name)/'template.docx';d=Document();d.add_paragraph('项目：{{项目}}');t=d.add_table(rows=2,cols=2);t.cell(0,0).text='字段';t.cell(0,1).text='值';d.save(path)
  data,_=export_template('|字段|值|\n|---|---|\n|项目|GSA02|','docx',path);d=Document(BytesIO(data));self.assertEqual(d.paragraphs[0].text,'项目：GSA02');self.assertEqual(d.tables[0].cell(1,1).text,'GSA02')
