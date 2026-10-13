import json, sqlite3, tempfile, unittest, os
from pathlib import Path
from contextlib import contextmanager
from unittest.mock import patch
from fastapi import FastAPI
from fastapi.testclient import TestClient
import role_config

class RoleConfigTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix="role-config-test-",dir=os.environ.get("PRESALES_TEST_TMP"))
        self.path=Path(self.tmp.name)/"db.sqlite"
        @contextmanager
        def db():
            c=sqlite3.connect(self.path,timeout=10);c.row_factory=sqlite3.Row
            try:yield c;c.commit()
            except:c.rollback();raise
            finally:c.close()
        self.db=db;self.account="software";self.files=Path(self.tmp.name)
        (self.files/"GSA02").mkdir()
        m={"manifest":[],"spec_pages":[],"schedule":{"total":100},"warnings":[],"phases":[{"id":"p"+str(i),"low":100,"high":100} for i in range(1,4)]}
        (self.files/"GSA02"/"materials.json").write_text(json.dumps(m))
        app=FastAPI();role_config.install(app,db,lambda:self.account,lambda:"now",self.files)
        self.client=TestClient(app);self.base="/api/role-config/software"
        self.env=patch.dict(os.environ,{"ZHIPU_API_KEY":"test-key","ZHIPU_MODEL":"test-model"});self.env.start()
        self.mock=patch.object(role_config,"_request_chat_stream",side_effect=lambda *a,**k:iter(["**真实接口模拟","输出**"]));self.mock.start()
        self.cfg=self.client.get(self.base).json()["versions"][0]
    def tearDown(self):
        self.mock.stop();self.env.stop();self.tmp.cleanup()
    def trial(self,task):
        payload={"config_id":self.cfg["id"],"task_id":task["id"],"project_id":"GSA02","question":"生成测试草稿","use_gsa_excerpt":True}
        if "quote_sheet" in task["tools"]:payload["quote_rows"]=[{"category":"软件授权费","name":"测试条目","quantity":2,"price":100}]
        if "server_sizing" in task["tools"]:payload["parameters"]={"scope":"phases","sampling_confirmed":True,"scope_confirmed":True,"assumptions_confirmed":True,"scope_note":"测试范围"}
        r=self.client.post(self.base+"/trials",json=payload);self.assertEqual(r.status_code,200,r.text)
        identifier=r.json()["id"];stream=self.client.get(self.base+"/trials/"+identifier+"/stream")
        values=[json.loads(line[6:]) for line in stream.text.splitlines() if line.startswith("data: ")]
        self.assertEqual(values[-1]["status"],"completed")
        return identifier,values[-1]
    def test_seed_and_permission(self):
        self.assertEqual(len(self.cfg["tasks"]),5)
        self.account="sales";self.assertEqual(self.client.get(self.base).status_code,403)
    def test_draft_immutability_and_conflict(self):
        tasks=self.cfg["tasks"];tasks[0]["name"]="更改入口"
        payload={"expected_revision":1,"tasks":tasks}
        self.assertEqual(self.client.post(self.base+"/drafts",json=payload).json()["revision"],2)
        self.assertEqual(self.client.post(self.base+"/drafts",json=payload).status_code,409)
        versions=self.client.get(self.base).json()["versions"]
        self.assertNotEqual(versions[-1]["tasks"][0]["name"],"更改入口")
    def test_tools_and_records(self):
        _,r=self.trial(next(t for t in self.cfg["tasks"] if t["id"]=="server"))
        tool=next(e for e in r["events"] if e["type"]=="tool")
        self.assertEqual(tool["output"]["points"]["high"],300)
        self.assertEqual(r["provider"],"zhipu")
        self.assertEqual(r["config_id"],self.cfg["id"])
    def test_handoff_uses_generation_config_after_project_upgrade(self):
        import workspace_router
        app=FastAPI()
        with self.db() as c:
            c.executescript('CREATE TABLE documents(id TEXT,project_id TEXT,title TEXT,version TEXT,sha256 TEXT); CREATE TABLE document_pages(document_id TEXT,page INT,text TEXT);')
            c.execute("UPDATE role_config_versions SET status='published' WHERE id=?",(self.cfg['id'],))
        workspace_router.install(app,self.db,lambda:self.account,lambda:'now',lambda *a:None)
        client=TestClient(app);base='/api/projects/A/workspaces/F5'
        with patch.object(workspace_router,'request_chat_stream',side_effect=lambda *a,**k:iter(['已知内容'])):
            r=client.post(base+'/ask-stream',json={'question':'功能清单','turn_id':'origin','task_id':'functions'})
            self.assertIn('done',r.text)
        tasks=json.loads(json.dumps(self.cfg['tasks']));next(t for t in tasks if t['id']=='functions')['recipient']='none'
        newer=self.client.post(self.base+'/drafts',json={'expected_revision':1,'tasks':tasks}).json()
        with self.db() as c:c.execute("UPDATE role_config_versions SET status='published' WHERE id=?",(newer['id'],))
        self.client.post(self.base+'/projects/A/bind',json={'config_id':newer['id']})
        r=client.post(base+'/artifacts/adopt',json={'source_id':'origin','title':'软件功能清单','content':'已知内容','conversation_id':'chat'})
        self.assertEqual(r.status_code,200,r.text)
        self.account='solution';messages=client.get(base+'/collaboration').json()
        self.assertEqual(len(messages),1);self.assertEqual(messages[0]['config_id'],self.cfg['id'])
        self.assertEqual(client.post(base+'/collaboration/'+messages[0]['id']+'/read').status_code,200)
        self.assertEqual(client.delete(base+'/collaboration/'+messages[0]['id']).status_code,200)
        self.assertEqual(client.get(base+'/collaboration').json(),[])
    def test_github_exports_only_configuration_and_reuses_pr(self):
        from types import SimpleNamespace
        with self.db() as c:c.execute("UPDATE role_config_versions SET status='published' WHERE id=?",(self.cfg['id'],))
        exported=[]
        def command(args,**kwargs):
            if args[0]=='git':return SimpleNamespace(returncode=0,stdout='feature/test')
            path=args[2]
            if path.endswith('/git/blobs'):
                exported.append(json.loads(kwargs['input'])['content']);value={'sha':'blob'}
            elif '/git/ref/heads/configs/' in path:return SimpleNamespace(returncode=1,stdout='')
            elif '/git/ref/' in path:value={'object':{'sha':'tip'}}
            elif '/git/commits/' in path:value={'tree':{'sha':'parent'}}
            elif path.endswith('/git/trees'):value={'sha':'tree'}
            elif path.endswith('/git/commits'):value={'sha':'commit'}
            elif path.endswith('/git/refs'):value={'ref':'branch'}
            elif '/pulls?' in path:value=[]
            elif path.endswith('/pulls'):value={'html_url':'https://github.com/Ponceau5/presales-workbench/pull/999'}
            else:raise AssertionError(path)
            return SimpleNamespace(returncode=0,stdout=json.dumps(value))
        with patch.object(role_config.shutil,'which',return_value='/bin/gh'),patch.object(role_config.subprocess,'run',side_effect=command):
            r=self.client.post(self.base+'/versions/'+self.cfg['id']+'/github');self.assertEqual(r.status_code,200,r.text)
            self.assertEqual(self.client.post(self.base+'/versions/'+self.cfg['id']+'/github').json(),r.json())
        value=json.loads(exported[0]);self.assertEqual(set(value),{'schema_version','role','revision','tasks'})
        self.assertNotIn('test-key',exported[0])
    def test_zhipu_balance_and_busy_are_distinct(self):
        import io
        from urllib.error import HTTPError
        import model_gateway
        cfg=model_gateway.ModelConfig('zhipu','glm','https://example.test','test-key')
        for code in ['1113','1305']:
            error=HTTPError('https://example.test',429,'error',{},io.BytesIO(json.dumps({'error':{'code':code}}).encode()))
            with patch.object(model_gateway,'urlopen',side_effect=error):
                with self.assertRaisesRegex(model_gateway.ModelError,code):list(model_gateway._request_chat_stream([],config=cfg))
    def test_word_excel_samples_and_exports(self):
        from io import BytesIO
        from docx import Document
        from openpyxl import load_workbook
        from file_formats import export_content,read_office,split_gap
        text='## 内部差距分析表\n| 条款 | 结论 |\n| --- | --- |\n| 1 | 未知 |\n\n## 客户澄清表\n| 原文 | 问题 |\n| --- | --- |\n| 要求 | 请确认 |'
        parts=split_gap(text);self.assertEqual(len(parts),2);self.assertNotIn('内部差距',parts[1]['content'])
        data,_=export_content(parts[1]['content'],'docx','客户澄清表')
        d=Document(BytesIO(data));self.assertEqual(d.tables[0].cell(1,1).text,'请确认')
        r=self.client.post(self.base+'/samples',data={'project_id':'GSA02','task_id':'gap','purpose':'template'},files={'file':('clarification.docx',data,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        self.assertEqual(r.status_code,200,r.text)
        with self.db() as c:
            chosen=role_config.load_samples(c,[r.json()['id']],'GSA02','gap')
            self.assertIn('仅参考结构',chosen[0]['usage'])
            from fastapi import HTTPException
            with self.assertRaises(HTTPException):role_config.load_samples(c,[r.json()['id']],'OTHER','gap')
        cost='| 产品 | 数量 | 单价 | 总价 |\n| --- | --- | --- | --- |\n| 软件 | 2 | 100 | 200 |\n| 硬件 | 未知 | 未知 | 未知 |'
        data,_=export_content(cost,'xlsx','成本单')
        book=load_workbook(BytesIO(data));sheet=book.active
        self.assertEqual(sheet['D3'].data_type,'f');self.assertIsNone(sheet['B4'].value)
    def test_cost_formulas_and_unknown_prices(self):
        rows=[role_config.QuoteLine(category='软件功能费',name='已知条目',quantity=2,price=100),role_config.QuoteLine(category='硬件费',name='未知规格')]
        path=self.files/'test.xlsx';result=role_config.quote_workbook(rows,path)
        self.assertEqual(result['known_subtotal'],'200.00');self.assertFalse(result['total_complete'])
        from openpyxl import load_workbook
        sheet=load_workbook(path).active
        self.assertEqual(sheet['I2'].data_type,'f');self.assertIsNone(sheet['H3'].value)
    def test_export_and_unvalidated_github_gate(self):
        result=self.client.get(self.base+"/versions/"+self.cfg["id"]+"/export")
        self.assertEqual(result.status_code,200)
        exported=result.json();self.assertEqual(exported["schema_version"],1)
        self.assertNotIn("ZHIPU_API_KEY",result.text)
        self.assertNotIn("sources",exported)
        self.assertEqual(self.client.post(self.base+"/versions/"+self.cfg["id"]+"/github").status_code,409)
    def test_cross_project_sources_rejected(self):
        r=self.client.post(self.base+"/trials",json={"config_id":self.cfg["id"],"task_id":"survey","project_id":"OTHER","question":"测试","use_gsa_excerpt":True})
        self.assertEqual(r.status_code,422)
    def test_publish_gate_and_project_pin(self):
        identifiers=[]
        first,_=self.trial(self.cfg["tasks"][0])
        self.assertEqual(self.client.post(self.base+"/publish",json={"config_id":self.cfg["id"],"expected_revision":1,"run_ids":[first]}).status_code,409)
        for task in self.cfg["tasks"]:
            identifier,_=self.trial(task);identifiers.append(identifier)
            self.assertEqual(self.client.post(self.base+"/trials/"+identifier+"/verdict",json={"accepted":True,"note":"人工测试确认"}).status_code,200)
        r=self.client.post(self.base+"/publish",json={"config_id":self.cfg["id"],"expected_revision":1,"run_ids":identifiers});self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(self.client.get("/api/projects/GSA02/software-config").json()["revision"],1)
        self.client.post(self.base+"/drafts",json={"expected_revision":1,"tasks":self.cfg["tasks"]})
        self.assertEqual(self.client.get("/api/projects/GSA02/software-config").json()["revision"],1)
        with self.db() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM skill_versions").fetchone()[0],15)
            self.assertEqual(c.execute("SELECT COUNT(*) FROM workflow_versions").fetchone()[0],15)

if __name__=="__main__":unittest.main()
