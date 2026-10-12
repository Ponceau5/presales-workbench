"""File inputs and readable Word/Excel exports; files are data, not executable code."""
from pathlib import Path
from io import BytesIO
import re

def read_office(path):
    suffix=Path(path).suffix.lower()
    if suffix=='.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        from docx.oxml.ns import qn
        d=Document(path);text=[]
        for element in d.element.body:
            if element.tag==qn('w:p'):
                value=Paragraph(element,d).text
                if value.strip():text.append(value)
            elif element.tag==qn('w:tbl'):
                text.extend(' | '.join(c.text for c in row.cells) for row in Table(element,d).rows)
        return ['\n'.join(text)]
    from openpyxl import load_workbook
    book=load_workbook(path,read_only=True,data_only=False)
    result=[]
    try:
        for sheet in book:
            if sheet.max_row>20000 or sheet.max_column>100:raise ValueError('工作表过大，请导入本次需要的部分')
            lines=['工作表：'+sheet.title]
            for row in sheet.iter_rows(values_only=True):
                if any(v is not None for v in row):lines.append(' | '.join('' if v is None else str(v) for v in row))
            result.append('\n'.join(lines))
    finally:book.close()
    return result

def blocks(text):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i]
        if '|' in line and i+1<len(lines) and re.match(r'^\s*\|?\s*:?-{3,}',lines[i+1]):
            table=[[v.strip() for v in line.strip().strip('|').split('|')]];i+=2
            while i<len(lines) and '|' in lines[i]:table.append([v.strip() for v in lines[i].strip().strip('|').split('|')]);i+=1
            yield ('table',table);continue
        yield ('text',line);i+=1

def export_content(text,kind,title):
    output=BytesIO()
    if kind=='docx':
        from docx import Document
        from docx.shared import Pt
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        d=Document();normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
        normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        d.add_heading(title,0)
        def write(p,value):
            for index,part in enumerate(re.split(r'\*\*(.*?)\*\*',value)):
                run=p.add_run(part);run.bold=bool(index%2)
        for kind_,value in blocks(text):
            if kind_=='table':
                cols=max(map(len,value));table=d.add_table(rows=0,cols=cols);table.style='Table Grid'
                for row in value:
                    cells=table.add_row().cells
                    for i,v in enumerate(row):write(cells[i].paragraphs[0],v)
                tr=table.rows[0]._tr.get_or_add_trPr();tr.append(OxmlElement('w:tblHeader'))
            elif value.strip():
                heading=re.match(r'^(#{1,6})\s+(.*)',value)
                if heading:write(d.add_heading(level=min(len(heading[1]),3)),heading[2])
                else:write(d.add_paragraph(),value.lstrip('- ') if value.startswith('- ') else value)
        d.save(output)
        return output.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if kind=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        book=Workbook();sheet=book.active;sheet.title='成果';sheet.append([title])
        for kind_,value in blocks(text):
            if kind_=='table':
                headers=[re.sub(r'\*\*','',v) for v in value[0]]
                qty=next((i for i,h in enumerate(headers) if h=='数量'),None);price=next((i for i,h in enumerate(headers) if h=='单价'),None);amount=next((i for i,h in enumerate(headers) if h in {'总价','金额','小计'}),None)
                for index,row in enumerate(value):
                    # Model text may resemble formulas; only the deterministic cost tool creates formulas.
                    safe=["'"+v if v.startswith(('=','+','-','@')) else re.sub(r'\*\*(.*?)\*\*',r'\1',v) for v in row]
                    sheet.append(safe)
                    if index and qty is not None and price is not None and amount is not None:
                        from openpyxl.utils import get_column_letter
                        n=sheet.max_row
                        for col in [qty,price]:
                            raw=re.sub(r'[，,￥¥元\s]','',row[col] if col<len(row) else '')
                            sheet.cell(n,col+1).value=float(raw) if re.fullmatch(r'\d+(?:\.\d+)?',raw) else None
                        q=get_column_letter(qty+1)+str(n);p=get_column_letter(price+1)+str(n)
                        sheet.cell(n,amount+1,f'=IF(OR({q}="",{p}=""),"",{q}*{p})')
                    if index==0:
                        for cell in sheet[sheet.max_row]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='218364')
            elif value.strip():sheet.append([re.sub(r'^#{1,6}\s+|\*\*','',value)])
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width=35
            for cell in col:cell.alignment=Alignment(wrap_text=True,vertical='top')
        sheet.freeze_panes='A2';book.save(output)
        return output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    raise ValueError('不支持的导出格式')


def split_gap(content):
    pattern=r'^#{1,3}\s*(内部差距分析表|客户澄清表)[^\n]*\n'
    found=list(re.finditer(pattern,content,re.M))
    if {m.group(1) for m in found}!={'内部差距分析表','客户澄清表'}:return []
    return [{'title':m.group(1),'content':content[m.start():found[i+1].start() if i+1<len(found) else len(content)].strip()} for i,m in enumerate(found)]
