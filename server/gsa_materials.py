"""导入指定项目原件的快照；来源是数据，不执行文档内指令。"""
from pathlib import Path
import hashlib
import json
import shutil
from collections import Counter
from openpyxl import load_workbook
import xlrd
from pypdf import PdfReader

ROOT = Path("/Users/a48750/Desktop/数据中心海外资料/项目相关文档/2026年/泰国GSA02数据中心项目")
VAULT = Path("/Users/a48750/Desktop/胡珊珊/CET/CET-Vault")
SKILL = Path("/Users/a48750/.openclaw/agents/main/agent/workshop-skills/scale-and-spec-estimation/SKILL.md")
MODEL = VAULT / "30-技术/服务器规格模型/服务器规格设计模型_v1.0.md"
NOTE = VAULT / "10-项目/进行中/泰国GSA02/泰国GSA02.md"
PHASES = [
 {"id":"p1","name":"一期","mw":12.5,"range":"3 个机房、4 台冷机、公共系统及平台基础","low":34000,"high":34000,"devices_low":990,"devices_high":990},
 {"id":"p2","name":"二期新增","mw":15,"range":"3 层 2 个机房、电力方舱、机房末端设备及部分剩余冷站","low":22000,"high":28000,"devices_low":600,"devices_high":750},
 {"id":"p3","name":"三期新增","mw":10.5,"range":"2 层 2 个机房、4 层 1 个机房、电力方舱及其余冷站","low":15000,"high":21000,"devices_low":450,"devices_high":650},
]

def import_materials(files_dir: Path) -> dict:
    snapshot = files_dir / "GSA02"
    snapshot.mkdir(exist_ok=True)
    manifest = []
    paths = {"spec":ROOT/"250626 Electrical MV And LV Specification Rev.00.pdf",
             "iolist":ROOT/"GSA02_IOLIST.xlsx", "supplement":ROOT/"增补清单.xlsx",
             "points":next(ROOT.glob("*software点位.xls")),
             "note":NOTE, "model":MODEL, "skill":SKILL,
             "phase_image":VAULT/"98-附件/泰国GSA02/GSA02-总体估算.png"}
    stored = {}
    for key, original in paths.items():
        digest = hashlib.sha256(original.read_bytes()).hexdigest()
        dest = snapshot / (key+"-"+digest[:12]+original.suffix)
        if not dest.exists(): shutil.copyfile(original, dest)
        stored[key] = dest
        manifest.append({"id":key,"title":original.name,"sha256":digest,"original_path":str(original),
                         "snapshot_name":dest.name})
    book = xlrd.open_workbook(str(stored["points"]))
    rows = []
    for sheet_name in ["EPMS (2)","BMS (2)"]:
        sheet = book.sheet_by_name(sheet_name)
        for col in range(2,sheet.ncols):
            count, coefficient = sheet.cell_value(1,col), sheet.cell_value(2,col)
            if isinstance(count,(int,float)) and isinstance(coefficient,(int,float)):
                rows.append({"system":sheet_name.split()[0],"device_type":str(sheet.cell_value(0,col)),
                             "quantity":count,"points_per_device":coefficient,
                             "points":count*coefficient,
                             "source":f"{paths['points'].name} · {sheet_name} · 行2/3 · 列{col+1}",
                             "warning":"单台 2700 点为手写估算，待实测复核" if str(sheet.cell_value(0,col))=="ECC" else ""})
    totals = {s:sum(r["points"] for r in rows if r["system"]==s) for s in ["EPMS","BMS"]}
    counts = Counter()
    workbook = load_workbook(stored["iolist"], data_only=True)
    for sheet in workbook:
        if sheet.title in {"Gateway","设备类型","FMS-PLC-IO-LIST"}: continue
        for row in sheet.iter_rows(min_row=2,values_only=True):
            if len(row)>7 and row[7] and str(row[7]).strip() not in {"Equipment Type","0"}:
                counts[str(row[7]).strip()] += 1
    reader = PdfReader(stored["spec"])
    # 本次软件演示选取 CMMS 两页，避免将整本 387 页无差别发送。
    pages = [{"page":i+1,"text":reader.pages[i].extract_text() or ""} for i in [220,221]]
    supplement = load_workbook(stored["supplement"],data_only=True).active
    historical = [{"row":i,"name":str(row[1]),"spec":str(row[2]),"quantity":row[5]}
                  for i,row in enumerate(supplement.values,1) if len(row)>5 and row[1] in {"DCOM Server","Server OS","DCOM Software"}]
    result = {"manifest":manifest,"schedule":{"rows":rows,"totals":totals,"total":sum(totals.values())},
              "iolist":{"typed_rows":sum(counts.values()),"by_type":dict(counts),
                        "warning":"按有设备类型的行计数，不等于确认的交付设备数；不得与点表数量直接相加。"},
              "phases":PHASES, "phase_source":"项目笔记《总体估算表》及 GSA02-总体估算.png，经验估算待审核",
              "spec_pages":pages, "historical":historical,
              "skill_text":stored["skill"].read_text(), "model_text":stored["model"].read_text(),
              "note_text":stored["note"].read_text(),
              "warnings":["点表、IOLIST、分期估算对应阶段尚未对齐。",
                          "楼栋级 BMS/EPMS 不是园区 DCOM 的实际接入量。",
                          "本项目已确认无 Kafka、无北向，历史备选方案不作为实施范围。",
                          "Server OS 4 套与 DCOM Server 2 台数量不一致。"]}
    (snapshot/"materials.json").write_text(json.dumps(result,ensure_ascii=False,indent=2))
    return result
